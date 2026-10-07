"""Read the supplied SOP without changing it; publish derived planning metadata only."""
import argparse
import csv
import hashlib
import json
import re
from collections import Counter, OrderedDict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/planning-v1.4'
CONFIG = ROOT / 'configs/scenarios-v1.4.json'
BENCHMARK = ROOT / 'tests/fixtures/full-scenario-benchmark-v1.4.json'

# Design assignments, not departments or permissions asserted by the source SOP.
DOMAINS = [
    ('订单与物流', '仓储物流；改址由主管核验', '订单/包裹状态、发货与揽收时间、轨迹、签收及拦截结果', '订单和物流只读；拦截写入另行获权', '核查→等待物流回执→回复；异常转主管', '物流核查结果或拦截终态已取得，客户结果已反馈'),
    ('包装破损与实物异常', '仓储物流；涉及安全交质量负责人', '商品明细、受损数量、包装状态、证据引用、使用影响、既往补发', '订单/售后只读；证据引用；仓储反馈', '补证→责任核查→人工方案→结果回传', '受损范围已确认，获批方案执行有回执且已反馈客户'),
    ('品质与食品安全', '质量负责人（新增职责）；主管', '批次、日期、开封与保存状态、安全信号、核验结论引用', '商品知识/质量证据；质量工单；售后只读', '安全分诊→必要补证→质量核验→主管方案', '质量负责人有结论，高风险处置与客户反馈完成'),
    ('口感与产品预期', '客服；运营；质量负责人按风险介入', 'SKU/版本、商品描述、使用方法、主观体验与客观异常区分', '版本化商品知识；订单只读', '知识解释→诉求确认→异常核查或人工方案', '问题解释有依据，仍有异议则保持处理中或升级'),
    ('错漏发与配货', '仓储物流；主管', '应发/实收明细、拆包与分包裹状态、数量差异、仓库复核', '订单/出库/包裹只读；仓库复核；售后', '排除分包裹→证据核验→人工方案→补发/退款回执', '数量差异有核查结论，补发或资金动作确认完成'),
    ('活动赠品与价格', '运营/活动负责人；财务按方案介入', '活动版本与有效期、订单资格、支付金额、赠品/券/奖品记录', '活动知识；订单/营销权益只读；通知', '核对活动资格→查兑现记录→人工补救→反馈', '活动兑现/不适用的依据明确，争议已交主管或处理完成'),
    ('退换退款与资金', '售后主管；财务；仓储物流', '订单与售后状态、退回验收、金额明细、消费者方案确认、审批记录', '订单/售后/退款进度只读；资金写入独立授权', '状态核实→方案审批→客户确认→执行回执→反馈', '退款到账/退货验收/换货履约按所选方案确认，不能以已申请结案'),
    ('发票会员与服务', '客服主管；财务或会员运营', '请求类型、资格/申请引用、当前处理人、服务等待和承诺记录', '工单/会员/开票状态只读；通知', '确定请求→查状态/资格→分派→回执→反馈', '开票/权益/服务事项有结果；客户未回复按挂起策略处理'),
    ('投诉与仓库系统协作', '主管分诊；运营/仓储/供应链/财务按事项接单', '争议或系统事件、各系统状态与时间、责任依据、审批、平台期限', '平台/ERP只读；部门任务与通知；写入另授权', '状态对齐→异常分诊→人工核实→执行协调→反馈', '平台期限已处理且所有相关任务有结果，资金与物流状态无冲突'),
    ('商品知识与渠道专项', '客服；品牌/运营；专项负责人', '平台/店铺、SKU与版本、有效知识/资质、问题事实、适用范围', '知识库；渠道规则；资质受控查阅；订单只读按需', '检索依据→核验适用范围→回复或专项任务', '回复有有效来源；未知参数/资料缺失转负责人且未虚构'),
]

DOMAIN_FACTS = {
    '订单与物流': ['order_ref', 'logistics_status', 'customer_request'],
    '包装破损与实物异常': ['product_ref', 'quantity', 'evidence', 'customer_request'],
    '品质与食品安全': ['product_ref', 'batch', 'opened_status', 'evidence', 'customer_request'],
    '口感与产品预期': ['product_ref', 'opened_status', 'customer_request'],
    '错漏发与配货': ['order_ref', 'product_ref', 'quantity', 'logistics_status', 'evidence'],
    '活动赠品与价格': ['order_ref', 'business_status', 'customer_request'],
    '退换退款与资金': ['order_ref', 'business_status', 'customer_request'],
    '发票会员与服务': ['order_ref', 'business_status', 'customer_request'],
    '投诉与仓库系统协作': ['order_ref', 'business_status', 'customer_request'],
    '商品知识与渠道专项': ['product_ref', 'business_status', 'customer_request'],
}

FACT_LABELS = {
    'order_ref': '订单引用（仅合成或脱敏）',
    'logistics_status': '物流/签收状态',
    'product_ref': '商品或SKU引用',
    'quantity': '应发、实收或受影响数量',
    'evidence': '证据是否已由获权人员核验',
    'batch': '批次/生产日期',
    'opened_status': '开封及保存状态',
    'customer_request': '客户当前诉求',
    'business_status': '活动/售后/服务系统状态',
}

def domain_index(i):
    for end, group in [(31, 0), (43, 1), (71, 2), (82, 3), (91, 4), (113, 5), (148, 6), (164, 7), (180, 8)]:
        if i <= end:
            return group
    # Late additions are assigned by business meaning, not their position in Excel.
    return {190:3,191:3,192:9,193:1,194:1,195:2,196:2,204:7,205:2,206:1,
            207:5,208:3,209:2,210:2,211:3,212:7,213:5,214:5,215:5,216:9,217:8,218:9}.get(i,9)

FACTS = [
    ('照片|拍摄|视频', '实物/场景证据引用（原始文件留在获权系统）'),
    ('订单号|订单编号', '订单引用（演示使用合成ID）'),
    ('面单|物流单号', '包裹/物流引用（不得留存未脱敏面单）'),
    ('数量|几件|几瓶', '应发、实收或受影响数量'),
    ('批次|批号|生产日期', '商品批次与生产日期'),
    ('开封|拆封', '是否开封及时间'),
    ('保存|储存', '保存条件'),
    ('受伤|物品受损|身体不适|腹泻|过敏', '安全/健康信号（不在演示保存健康原件）'),
    ('就医|诊断|病历|医疗', '受限凭证是否已由获权人员核验'),
    ('签收|收到商品|收到货', '签收与实收状态'),
    ('退款申请|退款状态|退款金额', '退款申请状态/金额依据'),
    ('下单.*时间|下单.*多久|购买时间', '下单/购买时间'),
    ('发票|税号|抬头', '开票需求及受控申请引用'),
    ('活动|优惠券|赠品|奖品|中奖', '活动规则及资格/兑现记录'),
]

P1 = {2,3,5,6,9,10,22,28,29,30,32,33,34,36,37,83,85,86,87,89,90,124,125,135,136,139,163,164,173,174,178}
KNOWN = {
    5: 'E39与E41的24/48小时分支不一致；需确认适用状态和规则版本',
    137: 'E1233称暂不支持换货，E1234支持同款换货；需明确店铺/款式范围',
    44: 'E395以开封后保存不当解释原因；不能仅凭已开封认定责任',
    178: 'E1605金额表述不清；E1602系统同步原因需以订单事件核实',
    121: '退款原因应匹配事实；E1088-E1089的话术需复核后使用',
}

def build(source):
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    if before != '8842d1acef3d5c25690e8e42385f963b1891db1ad46e3f72675b349cb1950022':
        raise ValueError('源表已变更，需重新评审场景映射；不能按旧顺序套用领域与批次。')
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb['售后问题汇总']
    groups = OrderedDict()
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        name = str(row[2] or '').strip()
        if name:
            groups.setdefault(name, []).append((n, str(row[3] or ''), str(row[4] or '')))
    records = []
    for i, (name, rows) in enumerate(groups.items(), 1):
        d, owner, data, connector, workflow, closure = DOMAINS[domain_index(i)]
        alltext = '\n'.join(t for _, _, t in rows)
        # Avoid counting a generic escalation template as evidence that the incoming case is high risk.
        safety = bool(re.search('安全|爆裂|发霉|变质|异物|絮状物|异味|发酵|涨袋|胀包|鼓包|过期|中毒|腹泻|过敏|身体不适|就医|治疗|赔付|赔偿|退款|退钱|投诉|平台介入|起诉|律师|改地址|改址|要挟', name))
        critical = safety or i in {1,17,18,119,141,142,143,146,151,153,217}
        risk = '高风险：人工主导' if critical else '条件风险：事实/权限校验'
        extracted = []
        locators = []
        for pattern, label in FACTS:
            hit = [n for n, stage, text in rows if ('首次' in stage or '索要' in stage or '确定诉求' in stage) and re.search(pattern,text)]
            if hit:
                extracted.append(label)
                locators.append(label + '：' + ','.join('E'+str(n) for n in hit))
        gates = []
        if critical: gates.append('该场景必须转人工；AI只做摘要/补证提示/路由建议')
        if re.search('退款|赔付|补偿|打款|理赔', alltext): gates.append('涉及资金的方案和执行必须人工审批；不能按话术自动兑付')
        gates.append('无资金/隐私变更的核查任务：商家明确预授权后可规则直派或人工批量批准；AI建议不等于授权')
        checks = ['商家/平台适用范围、审批权限、部门责任与时效尚未确认', '发生频率、每单耗时、实际错误率尚未提供']
        placeholders = [n for n,_,t in rows if re.search(r'XX|【\s*】',t)]
        if placeholders: checks.append('占位字段未定：'+','.join('E'+str(n) for n in placeholders))
        deadlines = sorted(set(re.findall(r'\d+(?:[-—~至]\d+)?\s*(?:个)?(?:工作日|小时|天)',alltext)))
        if deadlines: checks.append('原话术时效仅作候选；需明确计时起点/工作日/暂停/升级负责人')
        if '#分段#' in alltext: checks.append('含话术分段标记；补证要求需按场景筛选，不全部索要图片')
        if re.search('已处理完成|已加强|已优化|已反馈|已登记',alltext): checks.append('已完成/已反馈表述须有对应操作回执')
        if i in KNOWN: checks.insert(0,KNOWN[i])
        priority = 'P1 业务闭环' if i in P1 else ('P2 规则扩展' if i<=180 else 'P3 专项接入')
        reason = '可复用工单/库存基础，优先验证查询与催办价值；频率待验证' if i in P1 else '依赖商家专属规则或更复杂执行条件，基础完成后验证'
        if critical:
            reason = 'P0先覆盖风险识别/转人工；业务处置维持人工，领域自动化后置'
        baseline = ('可复用damage补证/审批样板，未逐场景实现' if 32<=i<=43 else
                    '可复用库存快照/替代候选/工单，未真实接入' if i in {95,139,174} else
                    '可复用工单与通知待发送箱，尚无此场景专用规则')
        response = ('仅可生成安全提示、必要补证和转人工草稿，人工确认发送' if critical else
                    '可生成解释/补证/已核实进度草稿；未来有效知识+新鲜事实+通道授权齐备才允许低风险自动回复')
        records.append({
            '场景ID':f'S{i:03d}', '原场景':name, '业务域（设计归类）':d,
            '原表定位':f'售后问题汇总!C{rows[0][0]}:E{rows[-1][0]}',
            '原表阶段': '；'.join(stage for _,stage,_ in rows),
            '原表明确输入（提取）':'；'.join(extracted) or '未提取到明确结构化输入，需场景评审',
            '输入依据单元格':'；'.join(locators) or '需场景评审',
            '系统需补数据（建议）':data,
            '风险与人工边界（建议）':risk,
            '回复边界（建议）':response,
            '审批方式（建议）':'；'.join(gates),
            '责任部门（待商家确认）':owner,
            '连接器需求（未验证）':connector,
            '处理流程（建议）':workflow,
            '结案条件（建议）':closure,
            '原话术时效（未批准）':'；'.join(deadlines) or '未提取到时效',
            '基础覆盖':'P0 分类/补证/人工兜底；不等于执行上线',
            '业务闭环批次（建议）':priority,
            '优先级理由':reason,
            '现有实现对照':baseline,
            '验收要点（建议）':f'{name}：完整事实走对应分支；缺关键输入只补证；风险升级使旧方案失效；无执行回执不结案；重复事件不重复执行',
            '待确认/疑点':'；'.join(checks),
            '业务频率':'未知；话术行数不能代表工单频率',
        })
    wb.close()
    assert len(records)==218 and len({r['场景ID'] for r in records})==218
    assert sum(len(v) for v in groups.values())==1969
    assert hashlib.sha256(source.read_bytes()).hexdigest()==before
    OUT.mkdir(parents=True, exist_ok=True)
    with (OUT/'scenario-matrix.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(records[0]))
        writer.writeheader(); writer.writerows(records)
    metadata={'source':'售后场景流程最终版.xlsx','sheet':'售后问题汇总','source_sha256':before,
              'scenario_count':len(records),'source_rows':1969,
              'domains':dict(Counter(r['业务域（设计归类）'] for r in records)),
              'priorities':dict(Counter(r['业务闭环批次（建议）'] for r in records)),
              'status':'设计初稿；部门、权限与字段映射尚需商家逐项确认',
              'extraction':'只从源表提取场景、阶段、输入关键词、时效和坐标；不保存员工名、客户问题原文或话术全文'}
    (OUT/'scenario-matrix.json').write_text(json.dumps({'meta':metadata,'rows':records},ensure_ascii=False,indent=2),encoding='utf-8')
    template=(ROOT/'scripts/full_scenario_plan.template.html').read_text(encoding='utf-8')
    payload=json.dumps({'meta':metadata,'rows':records},ensure_ascii=False).replace('<','\\u003c')
    result=template.replace('__PAYLOAD__',payload)
    (OUT/'index.html').write_text(result,encoding='utf-8')

    runtime_rows = []
    benchmark_rows = []
    for row in records:
        high_risk = row['风险与人工边界（建议）'].startswith('高风险')
        required = DOMAIN_FACTS[row['业务域（设计归类）']]
        runtime_rows.append({
            'id': row['场景ID'],
            'name': row['原场景'],
            'domain': row['业务域（设计归类）'],
            'source_locator': row['原表定位'],
            'source_fact_locator': row['输入依据单元格'],
            'required_facts': required,
            'required_fact_labels': [FACT_LABELS[key] for key in required],
            'risk_level': 'high' if high_risk else 'conditional',
            'owner_suggestion': row['责任部门（待商家确认）'],
            'priority': row['业务闭环批次（建议）'].split()[0],
            'review_status': 'business_review_required',
        })
        benchmark_rows.append({
            'case_id': 'TC-' + row['场景ID'],
            'scene_id': row['场景ID'],
            'synthetic_text': '合成测试：客户咨询“' + row['原场景'] + '”，请判断下一步。',
            'confirmed_facts': [],
            'expected_scene_id': row['场景ID'],
            'expected_route': 'human_triage' if high_risk else 'needs_evidence',
            'purpose': '验证218场景登记、精确标题识别及首轮安全分流；不代表自然语言泛化准确率',
        })
    runtime_meta = {
        'version': '1.4-p0-draft',
        'source': metadata['source'],
        'source_sha256': before,
        'scenario_count': 218,
        'review_status': 'business_review_required',
        'safe_scope': '仅保留场景名、来源坐标和设计元数据；不含员工名、客户原话或完整话术',
    }
    CONFIG.parent.mkdir(parents=True, exist_ok=True)
    BENCHMARK.parent.mkdir(parents=True, exist_ok=True)
    CONFIG.write_text(json.dumps({'meta': runtime_meta, 'fact_labels': FACT_LABELS, 'scenarios': runtime_rows}, ensure_ascii=False, indent=2), encoding='utf-8')
    BENCHMARK.write_text(json.dumps({'meta': runtime_meta, 'cases': benchmark_rows}, ensure_ascii=False, indent=2), encoding='utf-8')
    with (OUT/'full-scenario-test-cases.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(benchmark_rows[0]))
        writer.writeheader(); writer.writerows(benchmark_rows)
    checklist_rows = []
    for row in runtime_rows:
        checklist_rows.append({
            '场景ID': row['id'], '场景名称': row['name'], '业务域': row['domain'], '规则版本': runtime_meta['version'],
            '是否实际发生（待填写）': '', '近30天数量（待填写）': '', '当前人工步骤（待填写）': '', '当前使用系统（待填写）': '',
            '必须事实（建议）': '；'.join(row['required_fact_labels']), '责任部门（建议待确认）': row['owner_suggestion'],
            '是否可自动回复（待填写）': '', '是否可自动派核查任务（待填写）': '', '是否必须逐单审批（待填写）': '',
            '时效起点与目标（待填写）': '', '异常升级对象（待填写）': '', '确认人角色（待填写）': '',
            '确认日期（待填写）': '', '确认状态': '待业务确认', '备注': '',
        })
    with (OUT/'business-confirmation-checklist.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(checklist_rows[0]))
        writer.writeheader(); writer.writerows(checklist_rows)
    print(json.dumps(metadata,ensure_ascii=False,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('source',type=Path)
    build(parser.parse_args().source)
